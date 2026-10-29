"""Deterministic dataset import parsing for LabFlow.

Parsing is local and rule-based: pandas/openpyxl handle the file, and warnings are
derived from simple counts so the same upload always yields the same response.
CSV and XLSX share the same frame-validation, statistics, and preview pipeline.
"""

from __future__ import annotations

import csv
import io
from typing import Any
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.analysis import summary_stats

PREVIEW_ROWS = 10
MAX_UPLOAD_BYTES = 10 * 1024 * 1024
ALLOWED_CONTENT_TYPES = frozenset(
    {
        "text/csv",
        "application/csv",
        "application/vnd.ms-excel",
        "text/plain",
        "application/octet-stream",
    }
)
ALLOWED_XLSX_CONTENT_TYPES = frozenset(
    {
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/octet-stream",
    }
)


class DatasetImportError(ValueError):
    """Raised when an uploaded dataset cannot be accepted."""

    def __init__(self, message: str, status_code: int = 422) -> None:
        super().__init__(message)
        self.message = message
        self.status_code = status_code


def _json_safe(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, np.generic):
        value = value.item()
    if isinstance(value, float) and not np.isfinite(value):
        return None
    if isinstance(value, pd.Timestamp):
        return value.isoformat()
    return value


def _is_numeric_column(series: pd.Series) -> bool:
    return bool(pd.api.types.is_numeric_dtype(series)) and not bool(
        pd.api.types.is_bool_dtype(series)
    ) and int(series.notna().sum()) > 0


def _column_stats(series: pd.Series) -> dict[str, Any] | None:
    """Return summary statistics for a numeric column, or None if too few values."""
    values = pd.to_numeric(series, errors="coerce").to_numpy(dtype=float)
    finite = values[np.isfinite(values)]
    if finite.size < 2:
        return None
    stats = summary_stats(finite.tolist())
    return {
        "n": stats["n"],
        "missing": int(values.size - finite.size),
        "mean": stats["mean"],
        "std": stats["std"],
        "sem": stats["sem"],
    }


def _validate_field_counts(data: bytes) -> None:
    """Reject ragged records before pandas can silently shift or drop fields."""
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise DatasetImportError(f"CSV 格式错误：{exc}", status_code=422) from exc

    reader = csv.reader(io.StringIO(text, newline=""))
    expected: int | None = None
    try:
        for record in reader:
            if not record:
                continue
            if expected is None:
                expected = len(record)
                continue
            if len(record) != expected:
                raise DatasetImportError(
                    f"CSV 格式错误：第 {reader.line_num} 行有 {len(record)} 个字段，"
                    f"但表头有 {expected} 个",
                    status_code=422,
                )
    except csv.Error as exc:
        raise DatasetImportError(f"CSV 格式错误：{exc}", status_code=422) from exc


def _collect_warnings(frame: pd.DataFrame) -> list[str]:
    """Warnings every import reports, independent of format-specific notes."""
    warnings: list[str] = []

    missing_counts = frame.isna().sum()
    for column in frame.columns:
        missing = int(missing_counts[column])
        if missing > 0:
            warnings.append(f"列“{column}”有 {missing} 个缺失值")

    duplicate_rows = int(frame.duplicated().sum())
    if duplicate_rows > 0:
        warnings.append(f"数据集包含 {duplicate_rows} 个重复行")

    for column in frame.columns:
        series = frame[column]
        _describe_non_numeric(series, str(column), warnings)
        _describe_non_finite(series, str(column), warnings)

    return warnings


def _build_payload(
    frame: pd.DataFrame, format_label: str, warnings: list[str]
) -> dict[str, Any]:
    """Shared validation, warnings, statistics, and preview for every format."""
    if frame.empty:
        raise DatasetImportError(f"{format_label} 文件不含数据行", status_code=422)

    columns = [str(column) for column in frame.columns]

    warnings.extend(_collect_warnings(frame))

    numeric_columns = [
        str(column) for column in frame.columns if _is_numeric_column(frame[column])
    ]

    column_stats: dict[str, dict[str, Any]] = {}
    for column in numeric_columns:
        stats = _column_stats(frame[column])
        if stats is not None:
            column_stats[column] = stats

    head = frame.head(PREVIEW_ROWS)
    preview_frame = head.astype(object).where(pd.notna(head), None)
    preview = [
        {str(column): _json_safe(value) for column, value in row.items()}
        for row in preview_frame.to_dict(orient="records")
    ]

    return {
        "row_count": len(frame),
        "columns": columns,
        "numeric_columns": numeric_columns,
        "column_stats": column_stats,
        "preview": preview,
        "warnings": warnings,
    }


def load_csv_frame(data: bytes) -> tuple[pd.DataFrame, list[str]]:
    """Validate raw CSV bytes and return the parsed frame plus parser warnings."""
    if not data or not data.strip():
        raise DatasetImportError("CSV 文件为空", status_code=400)
    if len(data) > MAX_UPLOAD_BYTES:
        raise DatasetImportError(
            f"CSV 文件超过 {MAX_UPLOAD_BYTES // (1024 * 1024)} MB 上限",
            status_code=413,
        )

    _validate_field_counts(data)

    try:
        frame = pd.read_csv(io.BytesIO(data), encoding="utf-8-sig", index_col=False)
    except (pd.errors.EmptyDataError, pd.errors.ParserError, UnicodeDecodeError, ValueError) as exc:
        raise DatasetImportError(f"CSV 格式错误：{exc}", status_code=422) from exc

    return frame, []


def parse_csv_bytes(data: bytes) -> dict[str, Any]:
    """Parse raw CSV bytes into columns, numeric columns, preview, and warnings."""
    frame, warnings = load_csv_frame(data)
    return _build_payload(frame, "CSV", warnings)


def _normalize_cell(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, str) and not value.strip():
        return None
    return value


def _dedupe_columns(columns: list[str]) -> list[str]:
    originals = set(columns)
    used: set[str] = set()
    seen: dict[str, int] = {}
    deduped: list[str] = []
    for column in columns:
        if column not in seen:
            seen[column] = 0
            used.add(column)
            deduped.append(column)
            continue
        count = max(seen[column], 1)
        while True:
            candidate = f"{column}.{count}"
            count += 1
            if candidate not in originals and candidate not in used:
                break
        seen[column] = count
        used.add(candidate)
        deduped.append(candidate)
    return deduped


def _read_xlsx_rows(data: bytes) -> tuple[list[str], list[list[Any]]]:
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (
        BadZipFile,
        InvalidFileException,
        KeyError,
        OSError,
        ParseError,
        TypeError,
        ValueError,
    ) as exc:
        raise DatasetImportError(f"XLSX 格式错误：{exc}", status_code=422) from exc

    try:
        sheet_names = list(workbook.sheetnames)
        if not sheet_names:
            raise DatasetImportError("XLSX 工作簿不含工作表", status_code=422)
        sheet = workbook[sheet_names[0]]
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()
    return sheet_names, rows


def _xlsx_frame(sheet_name: str, rows: list[list[Any]]) -> pd.DataFrame:
    normalized: list[list[Any]] = []
    width = 0
    for row in rows:
        values = [_normalize_cell(value) for value in row]
        while values and values[-1] is None:
            values.pop()
        if not values:
            continue
        width = max(width, len(values))
        normalized.append(values)
    if not normalized or width == 0:
        raise DatasetImportError(
            f"工作表“{sheet_name}”不含数据行", status_code=422
        )

    header, *body = normalized
    padded_header = header + [None] * (width - len(header))
    columns = _dedupe_columns(
        [
            str(value) if value is not None else f"Unnamed: {index}"
            for index, value in enumerate(padded_header)
        ]
    )

    records = [row + [None] * (width - len(row)) for row in body]
    return pd.DataFrame(records, columns=columns)


def load_xlsx_frame(data: bytes) -> tuple[pd.DataFrame, list[str]]:
    """Validate raw XLSX bytes and return the first-sheet frame plus parser warnings."""
    if not data or not data.strip():
        raise DatasetImportError("XLSX 文件为空", status_code=400)
    if len(data) > MAX_UPLOAD_BYTES:
        raise DatasetImportError(
            f"XLSX 文件超过 {MAX_UPLOAD_BYTES // (1024 * 1024)} MB 上限",
            status_code=413,
        )

    sheet_names, rows = _read_xlsx_rows(data)
    frame = _xlsx_frame(sheet_names[0], rows)

    warnings: list[str] = []
    if len(sheet_names) > 1:
        warnings.append(
            f"工作簿包含 {len(sheet_names)} 个工作表；"
            f"已导入第一个工作表“{sheet_names[0]}”"
        )

    return frame, warnings


def parse_xlsx_bytes(data: bytes) -> dict[str, Any]:
    """Parse the first worksheet of raw XLSX bytes with the CSV response contract."""
    frame, warnings = load_xlsx_frame(data)
    return _build_payload(frame, "XLSX", warnings)


def _describe_non_numeric(series: pd.Series, column: str, warnings: list[str]) -> None:
    if _is_numeric_column(series) or pd.api.types.is_bool_dtype(series):
        return
    if not (pd.api.types.is_object_dtype(series) or pd.api.types.is_string_dtype(series)):
        return
    present = series[series.notna()]
    if present.empty:
        return
    coerced = pd.to_numeric(present, errors="coerce")
    non_numeric = int(coerced.isna().sum())
    if non_numeric > 0 and non_numeric < len(present):
        warnings.append(
            f"列“{column}”在数值数据中有 {non_numeric} 个非数值"
        )


def _describe_non_finite(series: pd.Series, column: str, warnings: list[str]) -> None:
    if not _is_numeric_column(series):
        return
    values = series.dropna().to_numpy(dtype=float)
    non_finite = int((~np.isfinite(values)).sum())
    if non_finite > 0:
        warnings.append(f"列“{column}”有 {non_finite} 个非有限值")
